# -*- coding: utf-8 -*-
"""实践任务三汇报 PPT 生成脚本（python-pptx）。"""
import os, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, "assets", "img")
SHOT = os.path.join(os.environ.get("TEMP", "."), "pptimg")
OUT = os.path.join(ROOT, "实践任务三-CodeBuddy初探-汇报.pptx")
TMP = os.path.join(os.environ.get("TEMP", "."), "pptcrop")
os.makedirs(TMP, exist_ok=True)

# ---------- 配色 ----------
INK = RGBColor(0x14, 0x21, 0x3D)      # 深墨蓝（深色页背景）
INK2 = RGBColor(0x1E, 0x2C, 0x4E)     # 卡片深底
PAPER = RGBColor(0xF7, 0xF5, 0xF0)    # 暖白（内容页背景）
CARD = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE4, 0xDF, 0xD5)
RED = RGBColor(0xC1, 0x27, 0x2D)      # 英雄红
RED_L = RGBColor(0xE0, 0x5A, 0x5F)
GOLD = RGBColor(0xE8, 0xB2, 0x3A)     # 满韵金
GOLD_L = RGBColor(0xF3, 0xD0, 0x8A)
DARK = RGBColor(0x1B, 0x24, 0x30)
SLATE = RGBColor(0x5A, 0x64, 0x73)
MIST = RGBColor(0x9A, 0xA5, 0xB8)     # 深色页上的弱化文字
FONT = "Microsoft YaHei"

W, H = 13.333, 7.5
ML = 0.72          # 左右边距
CW = W - 2 * ML    # 内容宽度 11.893


# ---------- 基础工具 ----------
def new_slide(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = INK if dark else PAPER
    return s


def tb(slide, x, y, w, h, text, size=14, bold=False, color=DARK,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=1.25, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = space
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
        r.font.name = FONT
    return box


def rect(slide, x, y, w, h, fill=None, line=None, rounded=0.0, lw=1.0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            shp.adjustments[0] = rounded
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(lw)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def accent(slide, x, y, h, w=0.075, color=RED):
    return rect(slide, x, y, w, h, fill=color)


def circle(slide, x, y, d, text, size=13, fill=RED, color=RGBColor(0xFF, 0xFF, 0xFF)):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = FONT
    return shp


def crop_fit(path, w, h, tag):
    """按目标宽高比裁剪，避免图片变形。"""
    im = Image.open(path).convert("RGB")
    tw, th = im.size
    target = w / h
    cur = tw / th
    if cur > target:
        nw = int(th * target)
        box = ((tw - nw) // 2, 0, (tw - nw) // 2 + nw, th)
    else:
        nh = int(tw / target)
        box = (0, (th - nh) // 2, tw, (th - nh) // 2 + nh)
    im = im.crop(box).resize((900, int(900 / target)), Image.LANCZOS)
    out = os.path.join(TMP, tag + ".png")
    im.save(out)
    return out


def pic(slide, path, x, y, w, h, tag, border=True):
    p = crop_fit(path, w, h, tag)
    s = slide.shapes.add_picture(p, Inches(x), Inches(y), Inches(w), Inches(h))
    if border:
        rect(slide, x, y, w, h, fill=None, line=LINE, rounded=0.03, lw=1.0)
    return s


def header(slide, kicker, title, dark=False, sub=None):
    c1 = GOLD if dark else RED
    c2 = RGBColor(0xFF, 0xFF, 0xFF) if dark else DARK
    c3 = MIST if dark else SLATE
    tb(slide, ML, 0.44, CW, 0.3, kicker, 11, True, c1)
    tb(slide, ML, 0.76, CW, 0.62, title, 27, True, c2)
    if sub:
        tb(slide, ML, 1.42, CW, 0.32, sub, 12, False, c3)


def footer(slide, idx, dark=False):
    c = MIST if dark else SLATE
    tb(slide, ML, 7.02, 6.0, 0.28,
       "实践任务三 · CodeBuddy 初探 ｜ 《人工智能导论》课程实践", 9, False, c)
    tb(slide, W - ML - 1.2, 7.02, 1.2, 0.28, "%02d" % idx, 9, True, c, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, dark=False, bar=True):
    rect(slide, x, y, w, h, fill=INK2 if dark else CARD,
         line=RGBColor(0x2C, 0x3B, 0x60) if dark else LINE, rounded=0.045)
    if bar:
        accent(slide, x, y + 0.22, h - 0.44)


# ---------- 数据 ----------
MBTI_URL = "https://btbct.github.io/mbti-test/"
HOME_URL = "https://btbct.github.io/siping-hometown/"

TYPE_COUNTS = [("INFP", 20), ("INTP", 18), ("ENFP", 17), ("INFJ", 15), ("ENFJ", 11),
               ("ENTP", 10), ("ISFJ", 10), ("ESFP", 9), ("INTJ", 9), ("ISTP", 8),
               ("ISFP", 7), ("ISTJ", 7), ("ENTJ", 6), ("ESFJ", 5), ("ESTP", 5), ("ESTJ", 3)]
DIM_AVG = [("E 外倾", 45.9, RED), ("I 内倾", 54.1, INK),
           ("S 实感", 42.3, RED), ("N 直觉", 57.7, INK),
           ("T 思考", 45.7, RED), ("F 情感", 54.3, INK),
           ("J 判断", 47.0, RED), ("P 知觉", 53.0, INK)]


def bar_chart(slide, x, y, w, h, cats, vals, colors=None, gap=60, label_size=9, title=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("s", vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = bool(title)
    if title:
        ch.chart_title.text_frame.text = title
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = DARK
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(label_size)
    dl.font.color.rgb = SLATE
    dl.font.name = FONT
    dl.number_format = '0.0"%"' if max(vals) <= 100 and isinstance(vals[0], float) else '0'
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    if colors:
        for i, c in enumerate(colors):
            try:
                pt = ser.points[i]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = c
            except Exception:
                pass
    else:
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RED
    try:
        ch.category_axis.tick_labels.font.size = Pt(10)
        ch.category_axis.tick_labels.font.name = FONT
        ch.category_axis.tick_labels.font.color.rgb = SLATE
        ch.category_axis.has_major_gridlines = False
        ch.value_axis.has_major_gridlines = False
        ch.value_axis.visible = False
        ch.value_axis.tick_labels.font.size = Pt(8)
    except Exception:
        pass
    return ch


# =========================================================
def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    # ---------- 01 封面 ----------
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, W, H, fill=INK)
    rect(s, 0, 0, 0.16, H, fill=RED)
    tb(s, 1.05, 1.35, 7.0, 0.36, "《人工智能导论》实践任务三 · CodeBuddy 初探", 13, True, GOLD)
    tb(s, 1.05, 1.95, 7.1, 1.5, "用 AI 编程助手\n从零做出两个网站", 40, True,
       RGBColor(0xFF, 0xFF, 0xFF), space=1.18)
    accent(s, 1.05, 3.62, 2.3, w=0.9, color=RED)
    tb(s, 1.05, 3.85, 6.9, 1.0,
       "① 网页版 MBTI 16 型人格测试与数据分析\n② “我的家乡”吉林四平文旅主题网站", 15, False,
       RGBColor(0xD8, 0xDF, 0xF0), space=1.5)
    tb(s, 1.05, 5.05, 6.9, 0.9,
       MBTI_URL + "\n" + HOME_URL, 11, False, GOLD_L, space=1.5)
    tb(s, 1.05, 6.25, 6.9, 0.7,
       "班级：____________　小组：____________\n姓名（学号）：__________________________",
       12, False, MIST, space=1.6)
    if os.path.exists(os.path.join(IMG, "hero.jpg")):
        pic(s, os.path.join(IMG, "hero.jpg"), 7.65, 1.05, 4.95, 5.4, "cover", border=False)
        rect(s, 7.65, 1.05, 4.95, 5.4, fill=None, line=GOLD, rounded=0.02, lw=1.2)

    # ---------- 02 目录 ----------
    s = new_slide(prs, dark=True)
    header(s, "CONTENTS", "汇报结构", dark=True, sub="按任务书的三项实践内容组织")
    items = [("01", "实践方案设计", "任务拆解 · AI Coding 闭环 · 环境设置"),
             ("02", "关键提示词实录", "需求表达 → 约束条件 → 反馈优化，共 10 条"),
             ("03", "两个网站成果", "MBTI 测试与数据分析 · 吉林四平文旅站"),
             ("04", "过程记录与复现", "5 类真实报错排查 · 6 步复现清单")]
    y = 2.15
    for no, t, d in items:
        card(s, ML, y, CW, 1.02, dark=True)
        circle(s, ML + 0.42, y + 0.26, 0.5, no, 13, GOLD, INK)
        tb(s, ML + 1.2, y + 0.24, 4.0, 0.4, t, 17, True, RGBColor(0xFF, 0xFF, 0xFF))
        tb(s, ML + 1.2, y + 0.62, 8.0, 0.32, d, 12, False, MIST)
        y += 1.18
    footer(s, 2, True)

    # ---------- 03 任务解读 ----------
    s = new_slide(prs)
    header(s, "01 / 实践方案", "任务书要什么，我怎么对应",
           sub="先拆评分点，再决定做什么、做到什么程度")
    cards = [
        ("网页版测试与分析", "MBTI 48 题测试：四个维度各 12 题，\n每题归属明确，结果有解释，\n并做了分布与倾向的可视化"),
        ("“我的家乡”文旅网站", "吉林四平专题站：5 大文旅分类、\n12 处景点、5 张文化名片、6 种美食、\n4 季玩法，手机端友好"),
        ("记录实施过程", "保留关键提示词原文、\nCodeBuddy 相关设置项，\n以及 5 类真实报错的排查记录"),
        ("成果展示", "本 PPT + 两个线上站点 + 数据图表，\n并给出 6 步可复现的操作清单"),
    ]
    pos = [(ML, 2.0), (ML + 5.95, 2.0), (ML, 4.35), (ML + 5.95, 4.35)]
    for (x, y), (t, d) in zip(pos, cards):
        card(s, x, y, 5.7, 2.15)
        tb(s, x + 0.42, y + 0.34, 4.8, 0.4, t, 16, True, DARK)
        tb(s, x + 0.42, y + 0.92, 4.9, 1.1, d, 12.5, False, SLATE, space=1.45)
    footer(s, 3)

    # ---------- 04 成果总览 ----------
    s = new_slide(prs)
    header(s, "OVERVIEW", "一页看完交付了什么",
           sub="两个独立仓库、两个线上站点，均可通过手机直接访问")
    stats = [("2", "个上线网站", "MBTI 测试站 · 四平文旅站"),
             ("48", "道测试题", "四个维度各 12 题"),
             ("32", "张实拍配图", "公开图源，压缩后 6.4MB"),
             ("20", "秒宣传短片", "站内图片自动合成"),
             ("9", "个内容板块", "分类/景点/人文/味道/季节…")]
    x = ML
    for n, lab, sub in stats:
        card(s, x, 2.05, 2.24, 2.5)
        tb(s, x + 0.28, 2.35, 1.8, 0.9, n, 44, True, RED, align=PP_ALIGN.CENTER)
        tb(s, x + 0.28, 3.35, 1.8, 0.35, lab, 13, True, DARK, align=PP_ALIGN.CENTER)
        tb(s, x + 0.18, 3.78, 1.95, 0.6, sub, 10.5, False, SLATE,
           align=PP_ALIGN.CENTER, space=1.35)
        x += 2.38
    tb(s, ML, 4.85, CW, 0.32, "线上地址", 13, True, DARK)
    rect(s, ML, 5.25, 5.85, 0.95, fill=CARD, line=LINE, rounded=0.06)
    accent(s, ML, 5.38, 0.69)
    tb(s, ML + 0.35, 5.42, 5.2, 0.28, "MBTI 16 型人格测试", 12, True, DARK)
    tb(s, ML + 0.35, 5.73, 5.2, 0.3, MBTI_URL, 10.5, False, RED)
    rect(s, ML + 6.15, 5.25, 5.74, 0.95, fill=CARD, line=LINE, rounded=0.06)
    accent(s, ML + 6.15, 5.38, 0.69)
    tb(s, ML + 6.5, 5.42, 5.2, 0.28, "我的家乡 · 吉林四平", 12, True, DARK)
    tb(s, ML + 6.5, 5.73, 5.2, 0.3, HOME_URL, 10.5, False, RED)
    footer(s, 4)

    # ---------- 05 AI Coding 闭环 ----------
    s = new_slide(prs)
    header(s, "01 / 实践方案", "AI Coding 的四步闭环",
           sub="需求表达 → 语义理解 → 代码生成 → 反馈优化，本次实践完整跑通了 5 轮")
    steps = [("需求表达", "把任务书拆成\n“两个网站 + 一份数据分析”，\n并给出参考站链接"),
             ("语义理解", "AI 读取参考站 HTML/CSS，\n归纳出“首页—答题—结果—分析”\n四视图结构"),
             ("代码生成", "按 index / css / js 分层生成，\n题库与文案独立成 data.js，\n改内容不用碰逻辑"),
             ("反馈优化", "截图自查 + 报错原文回灌，\n共修复 5 类问题\n（图显、编码、推送、证书、移动端）")]
    x = ML
    for i, (t, d) in enumerate(steps):
        card(s, x, 2.15, 2.72, 2.9)
        circle(s, x + 0.34, 2.45, 0.52, str(i + 1), 14)
        tb(s, x + 0.34, 3.12, 2.2, 0.4, t, 16, True, DARK)
        tb(s, x + 0.34, 3.62, 2.2, 1.2, d, 11.5, False, SLATE, space=1.45)
        if i < 3:
            tb(s, x + 2.66, 3.35, 0.35, 0.4, "›", 22, True, GOLD)
        x += 2.89
    rect(s, ML, 5.35, CW, 1.25, fill=INK, rounded=0.05)
    tb(s, ML + 0.42, 5.62, 10.8, 0.7,
       "体会：AI 生成的代码八成能直接用，剩下两成靠“把报错原文贴回去”解决。\n"
       "真正花时间的是把需求说清楚——说清楚之后，一轮就能拿到可用版本。",
       13, False, RGBColor(0xE2, 0xE8, 0xF8), space=1.5)
    footer(s, 5)

    # ---------- 06 环境准备与关键设置 ----------
    s = new_slide(prs)
    header(s, "02 / 实施过程", "环境准备与 CodeBuddy 关键设置",
           sub="这些设置项决定了后面能不能顺利跑通")
    left = [("工具与环境", True),
            ("• CodeBuddy：对话式改代码，直接读写项目文件", False),
            ("• 本地预览：tools/serve.ps1（PowerShell 静态服务器，无需 Node/Python）", False),
            ("• 版本与发布：Git + GitHub Pages，仓库根加 .nojekyll 原样发布", False),
            ("• 图片处理：PowerShell 调 .NET 重编码压缩（10.8MB → 6.4MB）", False)]
    right = [("关键设置", True),
             ("• git：http.sslBackend=schannel（绕开 OpenSSL 证书校验失败）", False),
             ("• PowerShell 脚本统一写 UTF-8 with BOM（5.1 缺 BOM 会中文乱码报错）", False),
             ("• tools/deploy.ps1：清洗用户名/仓库名输入、推送失败自动重试 3 次", False),
             ("• 移动端：390px 视口断点，首屏数据区改两列栅格", False)]
    for (x, items) in ((ML, left), (ML + 6.05, right)):
        card(s, x, 2.05, 5.84, 3.6)
        yy = 2.32
        for txt, is_h in items:
            if is_h:
                tb(s, x + 0.38, yy, 5.0, 0.38, txt, 16, True, DARK)
                yy += 0.62
            else:
                tb(s, x + 0.38, yy, 5.0, 0.62, txt, 12, False, SLATE, space=1.4)
                yy += 0.72
    rect(s, ML, 5.9, CW, 0.72, fill=RGBColor(0xFB, 0xF7, 0xEE), line=LINE, rounded=0.06)
    accent(s, ML, 6.02, 0.48, color=GOLD)
    tb(s, ML + 0.4, 6.1, 10.9, 0.4,
       "工作目录：c:\\Users\\ROG\\CodeBuddy\\（两个项目各一个独立 git 仓库，互不影响）",
       12, False, SLATE)
    footer(s, 6)

    # ---------- 07 提示词·需求表达 ----------
    s = new_slide(prs, dark=True)
    header(s, "02 / 关键提示词", "第一批：需求表达（做什么 + 长什么样）",
           dark=True, sub="给目标、给参考、给规模，AI 才不会跑偏")
    blocks = [
        ("提示词 ①", "帮我做一个网页版 MBTI 性格测试：48 道题，四个维度（E/I、S/N、T/F、J/P）分别计分，要有结果解释和数据分析；界面仿照 https://moyakura.github.io/ai-manju-survey/ 的暗色卡片风。"),
        ("提示词 ②", "再做一个“我的家乡”文旅网站，主题是吉林四平：要有合理的文旅项目分类、景点卡片配图、美食、季节玩法、行程规划，手机上也要好看。"),
    ]
    y = 2.2
    for tag, txt in blocks:
        rect(s, ML, y, CW, 1.55, fill=INK2, line=RGBColor(0x2C, 0x3B, 0x60), rounded=0.05)
        accent(s, ML, y + 0.28, 1.0, color=GOLD)
        tb(s, ML + 0.42, y + 0.28, 2.2, 0.35, tag, 13, True, GOLD)
        tb(s, ML + 0.42, y + 0.72, 10.9, 0.7, txt, 13.5, False,
           RGBColor(0xE8, 0xEC, 0xFF), space=1.4)
        y += 1.78
    tb(s, ML, 5.85, CW, 0.5,
       "要点：一句说清「交付物 + 规模 + 参考对象」，比追问十句更有效。", 12.5, False,
       GOLD_L)
    footer(s, 7, True)

    # ---------- 08 提示词·约束与生成 ----------
    s = new_slide(prs, dark=True)
    header(s, "02 / 关键提示词", "第二批：约束条件与代码组织",
           dark=True, sub="把边界和可维护性讲清楚，能省掉一轮返工")
    blocks = [
        ("提示词 ③", "题库用公开流传的 48 题版；注意第 31 题是反向题（A 选项对应情感），计分时别算反；每题顺序和 A/B 选项位置都要随机。"),
        ("提示词 ④", "把题库和 16 型人格文案单独放到 js/data.js，页面结构和交互逻辑分开，我以后只改数据文件就能换内容。"),
        ("提示词 ⑤", "网站要能在手机上打开，图片压缩一下，别让首屏加载太久。"),
    ]
    y = 2.05
    for tag, txt in blocks:
        rect(s, ML, y, CW, 1.12, fill=INK2, line=RGBColor(0x2C, 0x3B, 0x60), rounded=0.05)
        accent(s, ML, y + 0.22, 0.7, color=RED_L)
        tb(s, ML + 0.42, y + 0.18, 2.2, 0.3, tag, 12.5, True, RED_L)
        tb(s, ML + 0.42, y + 0.54, 10.9, 0.52, txt, 13, False,
           RGBColor(0xE8, 0xEC, 0xFF), space=1.4)
        y += 1.30
    tb(s, ML, 5.95, CW, 0.5,
       "效果：数据与逻辑分离后，换题库、改文案都不用碰代码，也不会引入新 bug。", 12.5, False,
       GOLD_L)
    footer(s, 8, True)

    # ---------- 09 提示词·反馈优化 ----------
    s = new_slide(prs, dark=True)
    header(s, "02 / 关键提示词", "第三批：反馈优化（把报错原文贴回去）",
           dark=True, sub="这是 AI Coding 里最省时间的一步")
    rows = [
        ("提示词 ⑥", "卡片里的图片显示不出来，只有深色块，帮我查原因。", "→ 占位层 z-index 盖住了 img"),
        ("提示词 ⑦", "手机上看首屏有横向溢出，文字被切掉了。", "→ 首屏数据区改两列栅格"),
        ("提示词 ⑧", "脚本报 UnexpectedToken；推送提示 repository not found / SSL certificate problem。",
         "→ BOM 编码 · 用户名清洗 · schannel"),
        ("提示词 ⑨", "网站要有视频素材，用站内图片生成一段 20 秒宣传短片并嵌进页面。",
         "→ ffmpeg 合成 + video 区块"),
    ]
    y = 2.0
    for tag, txt, res in rows:
        rect(s, ML, y, CW, 1.0, fill=INK2, line=RGBColor(0x2C, 0x3B, 0x60), rounded=0.05)
        accent(s, ML, y + 0.2, 0.64, color=GOLD)
        tb(s, ML + 0.4, y + 0.16, 1.5, 0.3, tag, 12, True, GOLD)
        tb(s, ML + 1.95, y + 0.16, 6.9, 0.62, txt, 12.5, False,
           RGBColor(0xE8, 0xEC, 0xFF), space=1.35)
        tb(s, ML + 9.0, y + 0.32, 2.7, 0.5, res, 11, True, GOLD_L, align=PP_ALIGN.RIGHT)
        y += 1.10
    tb(s, ML, 6.45, CW, 0.45,
       "提示词 ⑩（本次）：按实践任务三的要求做一份汇报 PPT，两个网站、提示词、实施过程、数据分析都放进去。",
       12.5, False, MIST)
    footer(s, 9, True)

    # ---------- 10 MBTI 网站 ----------
    s = new_slide(prs)
    header(s, "03 / 成果一", "MBTI 16 型人格测试站",
           sub="48 题 · 四维度计分 · 结果解释 · 本地数据统计")
    pts = [("题库与防偏差", "48 题 = 4 维度 × 12 题；题目顺序与 A/B 位置每次随机，\n避免“位置暗示答案”；第 31 题反向计分单独标注"),
           ("计分与结果", "按所选极计票 → 算百分比 → 每维度多数决 → 拼四字母；\n结果页给出维度条形图、优势/盲区、适合方向"),
           ("数据分析", "结果存入浏览器 localStorage，可看 16 型分布排行、\n四维度平均倾向，支持导出 CSV、载入示例数据、一键清空"),
           ("工程实现", "纯静态零依赖；index / css / js 分层，题库与文案独立；\n独立仓库 btbct/mbti-test，GitHub Pages 发布")]
    y = 2.05
    for t, d in pts:
        card(s, ML, y, 6.6, 1.12)
        tb(s, ML + 0.4, y + 0.2, 3.0, 0.34, t, 14, True, DARK)
        tb(s, ML + 0.4, y + 0.58, 5.9, 0.5, d, 11.5, False, SLATE, space=1.35)
        y += 1.24
    shot = os.path.join(SHOT, "mbti-home.png")
    if os.path.exists(shot):
        pic(s, shot, 7.5, 2.05, 5.1, 4.6, "mbhome")
        tb(s, 7.5, 6.72, 5.1, 0.3, "MBTI 测试站首页（暗色卡片风）", 10, False, SLATE,
           align=PP_ALIGN.CENTER)
    footer(s, 10)

    # ---------- 11 计分规则 ----------
    s = new_slide(prs)
    header(s, "03 / 成果一", "计分规则：四组偏好，各 12 题",
           sub="每题两选，选哪极就给哪极加一票，最后按多数决定字母")
    rows = [("维度", "题数", "两极", "在问什么"),
            ("能量来源 E / I", "12", "外倾 · 内倾", "注意力投向外部世界还是内心世界"),
            ("获取信息 S / N", "12", "实感 · 直觉", "相信事实细节，还是联想与可能性"),
            ("决策方式 T / F", "12", "思考 · 情感", "按逻辑准则判断，还是按人情价值取舍"),
            ("生活态度 J / P", "12", "判断 · 知觉", "偏好计划收束，还是保留弹性与可能")]
    y = 2.05
    for i, r in enumerate(rows):
        if i == 0:
            rect(s, ML, y, CW, 0.52, fill=INK, rounded=0.04)
            cols = [(ML + 0.4, 3.4), (ML + 4.0, 1.4), (ML + 5.6, 2.6), (ML + 8.5, 3.2)]
            for (cx, cw), txt in zip(cols, r):
                tb(s, cx, y + 0.14, cw, 0.3, txt, 12.5, True, RGBColor(0xFF, 0xFF, 0xFF))
            y += 0.62
        else:
            rect(s, ML, y, CW, 0.72, fill=CARD, line=LINE, rounded=0.04)
            vals = [(ML + 0.4, 3.4, True), (ML + 4.0, 1.4, False),
                    (ML + 5.6, 2.6, False), (ML + 8.5, 3.2, False)]
            for (cx, cw, b), txt in zip(vals, r):
                tb(s, cx, y + 0.22, cw, 0.32, txt, 12.5, b, DARK if b else SLATE)
            y += 0.82
    rect(s, ML, 5.95, CW, 0.85, fill=RGBColor(0xFB, 0xF7, 0xEE), line=LINE, rounded=0.06)
    accent(s, ML, 6.07, 0.61, color=GOLD)
    tb(s, ML + 0.42, 6.12, 10.9, 0.55,
       "自检：全部选 A 得 ESTJ（T 为 91.7%，说明第 31 题反向题确实没算反）；全部选 B 得 INFP。",
       12.5, False, SLATE, space=1.35)
    footer(s, 11)

    # ---------- 12 结果页 ----------
    s = new_slide(prs)
    header(s, "03 / 成果一", "结果页：不只是四个字母",
           sub="类型代号 + 中文名/族群 + 四维度百分比 + 优势盲区 + 适合方向")
    shot = os.path.join(SHOT, "mbti-result.png")
    if os.path.exists(shot):
        pic(s, shot, ML, 2.05, 5.35, 4.5, "mbres")
        tb(s, ML, 6.62, 5.35, 0.3, "结果页（示例：INFP 调停者）", 10, False, SLATE,
           align=PP_ALIGN.CENTER)
    items = [("① 类型与族群", "四字母代号 + 中文名，按四大族群（NT/NF/SJ/SP）配色"),
             ("② 维度倾向条", "E/I、S/N、T/F、J/P 各自百分比，落差一眼可见"),
             ("③ 优势与盲区", "双栏列出 3 条优势、2 条容易踩的坑"),
             ("④ 适合方向", "给出 3 类可考虑的职业/发展方向"),
             ("⑤ 留存与分享", "结果写入本地库，可复制结果文案、再测一次、查看分析")]
    y = 2.05
    for t, d in items:
        x = ML + 5.75
        card(s, x, y, 6.15, 0.9)
        tb(s, x + 0.38, y + 0.16, 2.6, 0.3, t, 13, True, DARK)
        tb(s, x + 0.38, y + 0.5, 5.5, 0.34, d, 11, False, SLATE)
        y += 1.0
    footer(s, 12)

    # ---------- 13 数据可视化：16 型分布 ----------
    s = new_slide(prs)
    header(s, "03 / 数据分析", "16 型分布：谁最多，谁最少",
           sub="示例数据 160 份答卷（用于演示分析流程，非真实普查）")
    cats = [c for c, _ in TYPE_COUNTS]
    vals = [v for _, v in TYPE_COUNTS]
    cols = []
    for c, _ in TYPE_COUNTS:
        fam = c[1] + c[2]
        cols.append({"NF": RGBColor(0x3E, 0x8E, 0x63), "NT": RGBColor(0x6A, 0x5A, 0xD8),
                     "SJ": RGBColor(0x3C, 0x7D, 0xB0), "SP": RGBColor(0xD9, 0x93, 0x2E)}.get(fam, RED))
    bar_chart(s, ML, 1.95, 7.5, 4.6, cats, vals, colors=cols, gap=45, label_size=8)
    tx = ML + 7.85
    card(s, tx, 1.95, 4.05, 4.6)
    tb(s, tx + 0.36, 2.2, 3.4, 0.35, "三条解读", 15, True, DARK)
    concl = ["样本 160 份，INFP(20)、INTP(18)、ENFP(17) 居前，\nNF 系（理想型）占比最高；",
             "ESTJ(3)、ESFJ(5)、ESTP(5) 最少，SJ 系（稳健型）\n在本样本中偏少；",
             "分布形态与同龄学生群体的常见结果一致：\n偏内向、偏直觉、偏情感、偏弹性。"]
    yy = 2.72
    for i, c in enumerate(concl):
        circle(s, tx + 0.36, yy, 0.34, str(i + 1), 11, RED)
        tb(s, tx + 0.82, yy + 0.02, 3.0, 0.9, c, 11.5, False, SLATE, space=1.4)
        yy += 1.24
    tb(s, ML, 6.65, 7.5, 0.3,
       "数据来源：测试站「载入示例数据」生成的 160 条记录，可导出 CSV 复核",
       10, False, SLATE)
    footer(s, 13)

    # ---------- 14 四维度倾向 ----------
    s = new_slide(prs)
    header(s, "03 / 数据分析", "四维度平均倾向：整体偏向哪一边",
           sub="同一批 160 份答卷，按四个维度分别统计平均百分比")
    cats = [c for c, _, _ in DIM_AVG]
    vals = [v for _, v, _ in DIM_AVG]
    cols = [c for _, _, c in DIM_AVG]
    bar_chart(s, ML, 1.95, 7.5, 4.3, cats, vals, colors=cols, gap=55, label_size=10)
    tx = ML + 7.85
    card(s, tx, 1.95, 4.05, 4.15)
    tb(s, tx + 0.36, 2.2, 3.4, 0.35, "结论", 15, True, DARK)
    concl = ["I 54.1% > E 45.9%：样本整体略偏内倾",
             "N 57.7% > S 42.3%：更关注可能性与意义",
             "F 54.3% > T 45.7%，P 53.0% > J 47.0%：偏情感、偏弹性",
             "四组落差均不足 16 个百分点：类型是倾向，不是标签"]
    yy = 2.66
    for i, c in enumerate(concl):
        circle(s, tx + 0.36, yy, 0.34, str(i + 1), 11, RED)
        tb(s, tx + 0.82, yy + 0.02, 3.0, 0.62, c, 11.5, False, SLATE, space=1.4)
        yy += 0.95
    rect(s, ML, 6.35, CW, 0.62, fill=INK, rounded=0.06)
    tb(s, ML + 0.4, 6.5, 11.0, 0.35,
       "说服力来自可复核：所有记录可一键导出 CSV，图表由同一批数据生成，换一批人重测即可复算。",
       12, False, RGBColor(0xE2, 0xE8, 0xF8))
    footer(s, 14)

    # ---------- 15 四平网站内容架构 ----------
    s = new_slide(prs)
    header(s, "03 / 成果二", "我的家乡 · 吉林四平文旅站",
           sub="9 个板块，从“怎么分类”到“怎么玩”一条线讲完")
    blocks = [("01 文旅分类", "红色记忆 / 满族风情 / 湖山自然 / 人文古迹 / 城市烟火"),
              ("02 好玩", "12 处景点与体验点，卡片看详情与实用贴士，♥ 加入心愿单"),
              ("03 人文", "5 张文化名片：让四平不只是“打过仗的城市”"),
              ("04 味道", "6 种四平美食：熏肉大饼、烧鸽子…"),
              ("05 季节", "春赏花 / 夏避暑 / 秋斑斓 / 冬戏雪，四季玩法"),
              ("06-09 其他", "行程规划 · 实用贴士 · 关于四平 · 资料来源可展开")]
    y = 2.0
    for i, (t, d) in enumerate(blocks):
        x = ML if i % 2 == 0 else ML + 6.05
        yy = y + (i // 2) * 1.22
        card(s, x, yy, 5.84, 1.06)
        tb(s, x + 0.4, yy + 0.18, 2.4, 0.3, t, 13.5, True, DARK)
        tb(s, x + 0.4, yy + 0.55, 5.1, 0.4, d, 11.5, False, SLATE)
    shot = os.path.join(SHOT, "home-home.png")
    if os.path.exists(shot):
        pic(s, shot, ML + 6.05, 5.7, 5.84, 1.2, "homecrop")
    footer(s, 15)

    # ---------- 16 交互与富媒体 ----------
    s = new_slide(prs)
    header(s, "03 / 成果二", "友好交互与富媒体素材",
           sub="文字 + 32 张配图 + 20 秒宣传短片，统一风格")
    left_items = [("分类筛选与详情", "点标签即筛选，点卡片弹窗看详情、贴士与来源"),
                  ("心愿单与行程规划", "♥ 收藏景点；按兴趣与天数自动生成路线与用餐建议"),
                  ("进度条与锚点导航", "顶部阅读进度条 + 9 个板块一键跳转"),
                  ("响应式适配", "390px 手机视口重排，首屏数据两列、卡片单列"),
                  ("资料可溯源", "“资料来源与数据出处”可展开，注明出处与许可"),
                  ("性能取舍", "配图统一压缩：10.8MB → 6.4MB，手机流量也能打开")]
    for i, (t, d) in enumerate(left_items):
        x = ML + (i % 2) * 3.6
        y = 2.0 + (i // 2) * 1.15
        card(s, x, y, 3.4, 1.0)
        tb(s, x + 0.36, y + 0.14, 2.8, 0.3, t, 12.5, True, DARK)
        tb(s, x + 0.36, y + 0.46, 2.8, 0.46, d, 10.5, False, SLATE, space=1.3)
    rx = ML + 7.35
    rw = 4.54
    p1 = os.path.join(IMG, "cat-red.jpg")
    if os.path.exists(p1):
        pic(s, p1, rx, 2.0, rw, 2.85, "spic0")
    for j, nm in enumerate(["cat-man.jpg", "food-01.jpg"]):
        p = os.path.join(IMG, nm)
        if os.path.exists(p):
            pic(s, p, rx + j * 2.34, 5.0, 2.2, 1.3, "spic%d" % (j + 1))
    rect(s, rx, 6.45, rw, 0.55, fill=INK, rounded=0.08)
    tb(s, rx, 6.58, rw, 0.3, "20 秒宣传短片 · 由站内图片自动合成", 11, True,
       RGBColor(0xE2, 0xE8, 0xF8), align=PP_ALIGN.CENTER)
    footer(s, 16)

    # ---------- 17 踩坑与解决 ----------
    s = new_slide(prs)
    header(s, "04 / 过程记录", "踩过的坑与排查结论",
           sub="每一条都来自本次实践的真实报错")
    rows = [("卡片图片只显示深色块", "占位渐变层 z-index 高于静态 img", "img 加 relative/z-index，加载成功后隐藏占位层"),
            ("一张配图下半截发灰", "图片下载被截断，非完整 JPEG", "校验 EOI 结尾标记，重新下载并全量复检 32 张"),
            ("PowerShell 脚本报 UnexpectedToken", "UTF-8 无 BOM 被 5.1 按 GBK 解码", "脚本统一写 UTF-8 with BOM，并做语法校验"),
            ("推送报 repository not found", "用户名填成了中文昵称（含全角括号）", "脚本自动清洗并校验用户名/仓库名格式"),
            ("推送报 SSL certificate problem", "Git 自带 OpenSSL 证书链校验失败", "改用 Windows 系统证书库（http.sslBackend=schannel）"),
            ("手机首屏横向溢出", "数据区一行挤爆容器", "首屏数据区改两列栅格，字号与内边距微调")]
    y = 2.0
    for a, b, c in rows:
        rect(s, ML, y, CW, 0.78, fill=CARD, line=LINE, rounded=0.04)
        accent(s, ML, y + 0.14, 0.5)
        tb(s, ML + 0.36, y + 0.24, 3.1, 0.3, a, 12.5, True, RED)
        tb(s, ML + 3.6, y + 0.24, 3.5, 0.3, b, 11.5, False, SLATE)
        tb(s, ML + 7.3, y + 0.24, 4.4, 0.5, c, 11.5, False, DARK, space=1.3)
        y += 0.83
    footer(s, 17)

    # ---------- 18 总结 ----------
    s = new_slide(prs, dark=True)
    rect(s, 0, 0, W, H, fill=INK)
    rect(s, 0, 0, 0.16, H, fill=RED)
    tb(s, 1.0, 0.75, 8.0, 0.36, "04 / 总结", 13, True, GOLD)
    tb(s, 1.0, 1.15, 8.0, 0.7, "六步即可复现", 30, True, RGBColor(0xFF, 0xFF, 0xFF))
    steps = ["安装 CodeBuddy，新建一个空工作目录",
             "按「需求表达 → 约束条件 → 反馈优化」的顺序提需求（提示词见第 7–9 页）",
             "本地 tools/serve.ps1 起服务，边看边改，截图自查",
             "git init → 提交 → tools/deploy.ps1 一键推送（失败会自动换证书后端重试）",
             "GitHub Pages 选 main / (root)，等 1–2 分钟",
             "手机打开网址验收，数据页可导出 CSV 复核"]
    y = 1.95
    for i, t in enumerate(steps):
        circle(s, 1.0, y, 0.42, str(i + 1), 12, GOLD, INK)
        tb(s, 1.6, y + 0.05, 6.7, 0.52, t, 12.5, False, RGBColor(0xE2, 0xE8, 0xF8), space=1.3)
        y += 0.66
    tb(s, 1.0, 5.78, 7.0, 0.62,
       "体会：AI 不是替你写代码，而是把「能不能把需求说清楚」\n直接变成生产力。",
       14, True, RGBColor(0xFF, 0xFF, 0xFF), space=1.45)
    tb(s, 1.0, 6.45, 7.0, 0.48, MBTI_URL + "\n" + HOME_URL, 10.5, False, GOLD_L, space=1.4)
    if os.path.exists(os.path.join(IMG, "cat-life.jpg")):
        pic(s, os.path.join(IMG, "cat-life.jpg"), 8.5, 1.6, 4.0, 4.2, "endpic", border=False)
        rect(s, 8.5, 1.6, 4.0, 4.2, fill=None, line=GOLD, rounded=0.02, lw=1.2)

    prs.save(OUT)
    print("SAVED:", OUT)
    print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))


if __name__ == "__main__":
    build()
